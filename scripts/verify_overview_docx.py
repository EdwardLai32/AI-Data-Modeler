"""Check the generated overview .docx opens cleanly and kept its content.

The renderer in ``build_overview_docx.py`` is hand-rolled, so the failure mode
worth guarding against is silent loss — a table or a section dropped without an
error. This reads the file back and asserts the load-bearing facts are present.
"""

from __future__ import annotations

import zipfile
from pathlib import Path

from docx import Document

TARGET = Path("AutoML Architect - Project Overview.docx")

# Figures and phrases that must survive; each is quoted in the document.
REQUIRED = [
    "772",
    "58,200",
    "0.9975",
    "0.791",
    "67,152",
    "min_frequency",
    "rationale",
    "Claude reasons",
    "AI-assisted",
    "out-of-fold",
]


def main() -> int:
    if not TARGET.exists():
        print(f"missing {TARGET}")
        return 1

    print(f"valid docx     : {zipfile.is_zipfile(TARGET)}")
    document = Document(str(TARGET))

    paragraphs = [p.text for p in document.paragraphs if p.text.strip()]
    print(f"paragraphs     : {len(paragraphs)}")
    shapes = [f"{len(t.rows)}x{len(t.columns)}" for t in document.tables]
    print(f"tables         : {len(document.tables)} -> {shapes}")

    headings = [p.text for p in document.paragraphs if p.style.name.startswith("Heading")]
    print(f"headings       : {len(headings)}")
    for heading in headings:
        print(f"   - {heading}")

    body = "\n".join(paragraphs)
    body += "\n".join(
        cell.text for table in document.tables for row in table.rows for cell in row.cells
    )

    print("\ncontent spot-check:")
    missing = []
    for needle in REQUIRED:
        present = needle in body
        if not present:
            missing.append(needle)
        print(f"   {'ok  ' if present else 'MISS'} {needle}")

    if missing:
        print(f"\nFAILED — {len(missing)} expected item(s) absent: {missing}")
        return 1
    print("\nall expected content present")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
